"""Per-suburb census "quirk" detector (ABS 2021 GCP, SAL geography).

For each mapped suburb we find the country-of-birth, language, and ancestry
groups that are *over-represented* versus the Greater Melbourne baseline,
measured as a Location Quotient (LQ = suburb's share / metro's share). LQ 3.0
means "three times the Melbourne rate". Paired with a percentile rank across
the mapped suburbs, this surfaces genuine character ("top 1% for Khmer") rather
than generic demographics.

Three lenses:
  G09  Country of Birth of Person      -> where residents themselves were born
  G13  Language Used at Home           -> language spoken
  G08  Ancestry by COB of Parents      -> heritage + a parents-born-overseas
                                          (1st-vs-2nd-gen) signal. NB: ABS does
                                          not publish *which country* a parent
                                          was born in, only overseas vs Aust.

Two confidence tiers (empirically tuned against the noise floor):
  core     : count >= 40 and LQ >= 2.0   (reliable community signal)
  emerging : count >= 25 and LQ >= 4.0   (rare-but-concentrated; lower n, so a
                                          steeper LQ is required to clear noise)
Cells with count < 20 are pure ABS-perturbation noise and never considered.

Run:
  uv run --with openpyxl --with pyshp python -m scrape.census Elsternwick
  uv run --with openpyxl --with pyshp python -m scrape.census --all
  uv run --with openpyxl --with pyshp python -m scrape.census --all --dry-run
"""
from __future__ import annotations

import csv
import glob
import io
import json
import os
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
PACK_DIR = ROOT / "data" / "2021_GCP_SAL_for_VIC_short-header"
DATA = PACK_DIR / "2021 Census GCP Suburbs and Localities for VIC"
META = PACK_DIR / "Metadata" / "Metadata_2021_GCP_DataPack_R1_R2.xlsx"
SHP_ZIP = ROOT / "data" / "raw" / "SAL_2021_AUST_GDA2020_SHP.zip"
BOUNDARIES = ROOT / "data" / "boundaries.geojson"
SUBURBS_JSON = ROOT / "data" / "suburbs.json"

CORE = dict(mincount=40, minlq=2.0)
EMERGING = dict(mincount=25, minlq=4.0)
TOP_N = 3  # per lens

# ABS classification names that read awkwardly in a UI -> friendlier label.
DISPLAY_RENAME = {
    "China excludes SARs and Taiwan": "China",
    "Korea Republic of South": "South Korea",
    "Hong Kong SAR of China": "Hong Kong",
    "Bosnia and Herzegovina": "Bosnia",
    "United States of America": "United States",
    "Other": "Other (unspecified)",
}

# Flag ISO codes (flag-icons set: lowercase alpha-2, plus gb-eng/sct/wls).
# Country of birth -> the country's own flag.
COUNTRY_ISO = {
    "Afghanistan": "af", "Bangladesh": "bd", "Bosnia": "ba", "Brazil": "br",
    "Cambodia": "kh", "Canada": "ca", "Chile": "cl", "China": "cn",
    "Croatia": "hr", "England": "gb-eng", "France": "fr", "Germany": "de",
    "Greece": "gr", "Hong Kong": "hk", "India": "in", "Indonesia": "id",
    "Iran": "ir", "Iraq": "iq", "Ireland": "ie", "Italy": "it", "Japan": "jp",
    "Lebanon": "lb", "Malaysia": "my", "Malta": "mt", "Mauritius": "mu",
    "Myanmar": "mm", "Nepal": "np", "Netherlands": "nl", "New Zealand": "nz",
    "North Macedonia": "mk", "Pakistan": "pk", "Papua New Guinea": "pg",
    "Poland": "pl", "Samoa": "ws", "Scotland": "gb-sct", "Singapore": "sg",
    "South Africa": "za", "South Korea": "kr", "Taiwan": "tw", "Thailand": "th",
    "Turkey": "tr", "United States": "us", "Vietnam": "vn", "Wales": "gb-wls",
    "Zimbabwe": "zw",
}
# Language -> a representative flag. A few (Arabic, Punjabi, Tamil) span several
# countries; we pick the origin dominant in Melbourne. None = show no flag.
LANG_ISO = {
    "Afrikaans": "za", "Arabic": "lb", "Australian Indigenous Languages": "au",
    "Bengali": "bd", "Cantonese": "hk", "Croatian": "hr", "French": "fr",
    "German": "de", "Greek": "gr", "Gujarati": "in", "Hindi": "in",
    "Indonesian": "id", "Italian": "it", "Japanese": "jp", "Khmer": "kh",
    "Korean": "kr", "Macedonian": "mk", "Malayalam": "in", "Mandarin": "cn",
    "Nepali": "np", "Persian excluding Dari": "ir", "Polish": "pl",
    "Portuguese": "pt", "Punjabi": "in", "Russian": "ru", "Samoan": "ws",
    "Serbian": "rs", "Sinhalese": "lk", "Spanish": "es", "Tamil": "lk",
    "Thai": "th", "Turkish": "tr", "Urdu": "pk", "Vietnamese": "vn",
    "Other (unspecified)": None,
}

# Language-family rollups (sums of their child languages) — exclude to avoid
# double counting. Their short codes end in `_Tot_Tot`.
_FAMILY_PREFIXES = (
    "Southeast Asian Austronesian languages ",
    "Indo Aryan languages ",
    "Chinese languages ",
    "Eastern European languages ",
    "Southwest and Central Asian languages ",
    "Dravidian languages ",
    "African languages ",
    "Other languages ",
)


# --------------------------------------------------------------------------- #
# metadata + geography
# --------------------------------------------------------------------------- #
def load_longmap() -> dict[str, str]:
    """short column code -> long human description (from the metadata workbook)."""
    import openpyxl

    wb = openpyxl.load_workbook(META, read_only=True)
    ws = wb["Cell Descriptors Information"]
    rows = list(ws.iter_rows(values_only=True))[11:]
    return {r[1]: r[2] for r in rows if r[1]}


def load_geo() -> tuple[dict[str, str], set[str]]:
    """Return (display_name -> 'SAL<code>', set of Greater-Melbourne SAL codes).

    Greater Melbourne is approximated as every VIC SAL whose centroid falls in
    the bounding box of the suburbs we actually map — a fair, self-contained
    denominator that includes metro suburbs we didn't profile while excluding
    regional Victoria.
    """
    import shapefile  # pyshp

    z = zipfile.ZipFile(SHP_ZIP)
    base = next(n[:-4] for n in z.namelist() if n.endswith(".shp"))
    r = shapefile.Reader(
        shp=io.BytesIO(z.read(base + ".shp")), dbf=io.BytesIO(z.read(base + ".dbf"))
    )
    name_to_code: dict[str, str] = {}
    centroid: dict[str, tuple[float, float]] = {}
    for shp, rec in zip(r.shapes(), r.records()):
        d = rec.as_dict()
        if d["STE_NAME21"] != "Victoria":
            continue
        bb = getattr(shp, "bbox", None)
        if not bb or not shp.points:
            continue
        name_to_code[d["SAL_NAME21"]] = d["SAL_CODE21"]
        centroid[d["SAL_CODE21"]] = ((bb[0] + bb[2]) / 2, (bb[1] + bb[3]) / 2)

    geo = json.loads(BOUNDARIES.read_text(encoding="utf-8"))
    display_to_code: dict[str, str] = {}
    mapped_codes: list[str] = []
    for feat in geo["features"]:
        p = feat["properties"]
        sal_name, display = p["sal_name"], p["suburb"]
        code = name_to_code.get(sal_name)
        if not code:
            print(f"  WARN: no SAL code for {sal_name!r}", file=sys.stderr)
            continue
        display_to_code[display] = "SAL" + code
        mapped_codes.append(code)

    xs = [centroid[c][0] for c in mapped_codes]
    ys = [centroid[c][1] for c in mapped_codes]
    x0, x1 = min(xs) - 0.05, max(xs) + 0.05
    y0, y1 = min(ys) - 0.05, max(ys) + 0.05
    metro = {
        "SAL" + c
        for c, (cx, cy) in centroid.items()
        if x0 <= cx <= x1 and y0 <= cy <= y1
    }
    return display_to_code, metro


# --------------------------------------------------------------------------- #
# table loading + label cleaning
# --------------------------------------------------------------------------- #
def load_table(glob_pat: str, keep: set[str]) -> dict[str, dict[str, str]]:
    by: dict[str, dict[str, str]] = {}
    for path in sorted(glob.glob(str(DATA / glob_pat))):
        with open(path, newline="", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                sal = row["SAL_CODE_2021"]
                if sal in keep:
                    by.setdefault(sal, {}).update(row)
    return by


def _strip_family(s: str) -> str:
    for fam in _FAMILY_PREFIXES:
        s = s.replace(fam, "")
    return s.strip()


def make_pickers(longmap: dict[str, str]):
    """Return {lens: picker}. Each picker maps a column code -> display label
    (or None to skip), with all known ABS-label artefacts fixed."""

    def language(k: str):
        if not (k.startswith("POL_") and k.endswith("_Tot")):
            return None
        if k == "POL_Tot_Tot" or k.endswith("_UOLSE_Tot") or k.endswith("_Tot_Tot"):
            return None  # grand total, English-subtotal, or family rollup
        raw = longmap.get(k, k)
        lab = raw.replace("PERSONS_Uses_other_language_", "").replace("_Total", "")
        return _strip_family(lab.replace("_", " "))

    def country(k: str):
        if not (k.startswith("P_") and k.endswith("_Tot")):
            return None
        if k in ("P_Tot_Tot", "P_Australia_Tot", "P_COB_NS_Tot", "P_Elsewhere_Tot"):
            return None
        raw = longmap.get(k, k)
        # Afghanistan's total is uniquely '..._Age_Total' -> drop the stray Age.
        lab = raw.replace("PERSONS_", "").replace("_Age_Total", "").replace("_Total", "")
        return lab.replace("_", " ").strip()

    def ancestry(k: str):
        if not k.endswith("_Tot_resp"):
            return None
        a = k[: -len("_Tot_resp")]
        if a in ("Tot", "Tot_P", "Ancestry_NS"):
            return None  # grand-total rows / not-stated
        if a.endswith(("_BP_B_OS", "_FO_B_OS", "_MO_B_OS", "_BP_B_Aus", "_BP_NS")):
            return None  # parent-birthplace breakdown sub-columns
        label = a.replace("Aust_Abor", "Australian Aboriginal").replace("Aust", "Australian")
        return label.replace("_", " ").strip()

    return {"language": language, "birthplace": country, "ancestry": ancestry}


def display_name(label: str) -> str:
    return DISPLAY_RENAME.get(label, label)


def _attach_iso(quirk: dict, lens: str) -> None:
    """Add a flag ISO code to language/birthplace quirks (ancestry gets none)."""
    if lens == "language":
        quirk["iso"] = LANG_ISO.get(quirk["group"])
    elif lens == "birthplace":
        quirk["iso"] = COUNTRY_ISO.get(quirk["group"])


# --------------------------------------------------------------------------- #
# LQ engine
# --------------------------------------------------------------------------- #
def build_groups(by: dict, picker) -> dict[str, dict[str, int]]:
    rows: dict[str, dict[str, int]] = {}
    for sal, d in by.items():
        g: dict[str, int] = {}
        for k, v in d.items():
            lab = picker(k)
            if lab and v.lstrip("-").isdigit():
                g[lab] = g.get(lab, 0) + int(v)
        rows[sal] = g
    return rows


def baseline_share(rows: dict, totals: dict) -> dict[str, float]:
    grp: dict[str, int] = {}
    base_tot = 0
    for sal, g in rows.items():
        base_tot += totals[sal]
        for k, v in g.items():
            grp[k] = grp.get(k, 0) + v
    return {k: v / base_tot for k, v in grp.items()}


def quirks_for(
    sal: str, rows: dict, base: dict, totals: dict, mapped: set[str]
) -> list[dict]:
    """All groups for one suburb that clear either tier, richest-first."""
    out = []
    for k, count in rows.get(sal, {}).items():
        if count < EMERGING["mincount"]:
            continue
        share = count / totals[sal]
        bs = base.get(k, 1e-9)
        lq = share / bs if bs else 0.0
        core = count >= CORE["mincount"] and lq >= CORE["minlq"]
        emerg = count >= EMERGING["mincount"] and lq >= EMERGING["minlq"]
        if not (core or emerg):
            continue
        shares = sorted(rows[s].get(k, 0) / totals[s] for s in mapped if s in rows)
        pct = sum(1 for x in shares if x <= share) / len(shares) * 100
        out.append(
            {
                "group": display_name(k),
                "lq": round(lq, 1),
                "percentile": round(pct),
                "top_pct": max(1, round(100 - pct)),
                "count": count,
                "share_pct": round(100 * share, 1),
                "baseline_pct": round(100 * bs, 2),
                "tier": "core" if core else "emerging",
            }
        )
    out.sort(key=lambda q: q["lq"], reverse=True)
    return out


def _lens_phrase(lens: str, group: str) -> str:
    if lens == "language":
        return f"{group} spoken at home"
    if lens == "birthplace":
        return f"{group}-born residents"
    return f"{group} ancestry"


def make_headline(per_lens: dict[str, list[dict]]) -> str | None:
    best, best_lens = None, None
    for lens, items in per_lens.items():
        for q in items:
            if q["tier"] == "core" and (best is None or q["lq"] > best["lq"]):
                best, best_lens = q, lens
    if not best:
        return None
    what = _lens_phrase(best_lens, best["group"])
    return (
        f"Top {best['top_pct']}% of Melbourne for {what} "
        f"({best['lq']:.1f}× the metro rate)"
    )


# --------------------------------------------------------------------------- #
# orchestration
# --------------------------------------------------------------------------- #
class Census:
    def __init__(self):
        longmap = load_longmap()
        self.display_to_code, self.metro = load_geo()
        self.pickers = make_pickers(longmap)

        g09 = load_table("2021Census_G09*_VIC_SAL.csv", self.metro)
        g13 = load_table("2021Census_G13*_VIC_SAL.csv", self.metro)
        g08 = load_table("2021Census_G08_VIC_SAL.csv", self.metro)
        self.totals = {sal: int(d["P_Tot_Tot"]) for sal, d in g09.items()}
        self.g08_raw = g08
        self.g09_raw = g09

        self.mapped = set(self.display_to_code.values())
        self.lenses = {}
        for lens, raw in (("language", g13), ("birthplace", g09), ("ancestry", g08)):
            rows = build_groups(raw, self.pickers[lens])
            self.lenses[lens] = (rows, baseline_share(rows, self.totals))

    def _summary_stats(self, sal: str, g09_row: dict) -> dict:
        total = int(g09_row["P_Tot_Tot"])
        aus = int(g09_row.get("P_Australia_Tot", 0))
        ns = int(g09_row.get("P_COB_NS_Tot", 0))
        overseas = total - aus - ns
        stats = {
            "population": total,
            "born_overseas_pct": round(100 * overseas / total, 1) if total else None,
        }
        a = self.g08_raw.get(sal, {})
        resp = int(a.get("Tot_P_Tot_resp", 0))
        if resp:
            bpo = int(a.get("Tot_P_BP_B_OS", 0))
            one = bpo + int(a.get("Tot_P_FO_B_OS", 0)) + int(a.get("Tot_P_MO_B_OS", 0))
            stats["both_parents_overseas_pct"] = round(100 * bpo / resp, 1)
            stats["at_least_one_parent_overseas_pct"] = round(100 * one / resp, 1)
        return stats

    def block_for(self, display: str) -> dict | None:
        sal = self.display_to_code.get(display)
        if not sal or sal not in self.totals:
            return None
        block = self._summary_stats(sal, self.g09_raw[sal])
        per_lens, emerging = {}, []
        for lens, (rows, base) in self.lenses.items():
            qs = quirks_for(sal, rows, base, self.totals, self.mapped)
            for q in qs:
                _attach_iso(q, lens)
            core = [q for q in qs if q["tier"] == "core"][:TOP_N]
            per_lens[lens] = core
            for q in qs:
                if q["tier"] == "emerging" and q not in core:
                    emerging.append({**q, "lens": lens})
        block.update(per_lens)
        block["emerging"] = sorted(emerging, key=lambda q: q["lq"], reverse=True)[:3]
        block["headline"] = make_headline(per_lens)
        return block


def main(argv: list[str]) -> int:
    args = [a for a in argv if not a.startswith("--")]
    flags = {a for a in argv if a.startswith("--")}
    census = Census()

    if "--all" in flags:
        suburbs = json.loads(SUBURBS_JSON.read_text(encoding="utf-8"))
        done, missing = 0, []
        for name in suburbs:
            block = census.block_for(name)
            if block is None:
                missing.append(name)
                continue
            suburbs[name]["census"] = block
            done += 1
        print(f"census: filled {done}/{len(suburbs)} suburbs", file=sys.stderr)
        if missing:
            print(f"  no census data for: {missing}", file=sys.stderr)
        if "--dry-run" in flags:
            sample = next(iter(suburbs))
            print(json.dumps(suburbs[sample]["census"], indent=2, ensure_ascii=False))
            print("(dry-run — suburbs.json not written)", file=sys.stderr)
        else:
            SUBURBS_JSON.write_text(
                json.dumps(suburbs, indent=2, ensure_ascii=False), encoding="utf-8"
            )
            print(f"wrote {SUBURBS_JSON}", file=sys.stderr)
        return 0

    if not args:
        print(__doc__)
        return 1
    block = census.block_for(args[0])
    if block is None:
        print(f"no census data for {args[0]!r}", file=sys.stderr)
        return 1
    print(json.dumps(block, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
